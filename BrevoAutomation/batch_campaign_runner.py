import os
import re
import json
import logging
import datetime
import openpyxl
from pathlib import Path

try:
    from BrevoAutomation.brevorequester import BrevoRequester
except ImportError:
    from brevorequester import BrevoRequester

SCRIPT_DIR = Path(__file__).parent
EXCEL_FILE = SCRIPT_DIR / "samlet.xlsx"
STATE_FILE = SCRIPT_DIR / "batch_state.json"
BATCH_SIZE = 1448  # 1448 + 2 pinned = 1450 total (50 under Brevo's 1500 limit)
TARGET_LIST_ID = 4

PINNED_CONTACTS = [
    {"email": "jeppebs02@gmail.com", "attributes": {"FIRSTNAME": ""}, "listIds": [TARGET_LIST_ID], "updateEnabled": True},
    {"email": "soerensenmorten@gmail.com", "attributes": {"FIRSTNAME": ""}, "listIds": [TARGET_LIST_ID], "updateEnabled": True},
]
PINNED_EMAILS = {c["email"].lower() for c in PINNED_CONTACTS}


def setup_logging():
    log_filename = SCRIPT_DIR / f"script_log_batch_{datetime.datetime.now():%Y%m%d_%H%M%S}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[
            logging.FileHandler(log_filename),
            logging.StreamHandler(),
        ],
    )
    logging.info(f"Log file: {log_filename}")


def load_state():
    if STATE_FILE.exists():
        with open(STATE_FILE, "r") as f:
            state = json.load(f)
        logging.info(f"Loaded state: offset={state.get('offset', 0)}, total={state.get('total_contacts', '?')}")
        return state
    logging.info("No state file found — starting from offset 0")
    return {"offset": 0, "total_contacts": 0}


def save_state(offset, total):
    state = {
        "offset": offset,
        "last_run": datetime.datetime.now().isoformat(),
        "total_contacts": total,
    }
    with open(STATE_FILE, "w") as f:
        json.dump(state, f, indent=2)
    logging.info(f"State saved: offset={offset}, total={total}")


def parse_samlet_xlsx(filepath, target_list_id):
    """Parse samlet.xlsx and return a list of Brevo contact dicts. Pinned emails are excluded."""
    contacts = []
    try:
        wb = openpyxl.load_workbook(filename=str(filepath), data_only=True)
        sheet = wb.active
    except FileNotFoundError:
        logging.error(f"Excel file not found: {filepath}")
        return []
    except Exception as e:
        logging.exception(f"Error loading workbook {filepath}: {e}")
        return []

    header_row_index = -1
    email_col_index = -1
    fname_col_index = -1
    MAX_HEADER_SEARCH_ROWS = 5

    for r_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=MAX_HEADER_SEARCH_ROWS, values_only=True)):
        row_values = [str(v).strip() if v is not None else "" for v in row]
        lower = [v.lower() for v in row_values]
        # Accept either column naming convention
        email_candidates = ["e-mail", "mail"]
        name_candidates = ["navn", "fornavn"]
        email_hit = next((c for c in email_candidates if c in lower), None)
        name_hit = next((c for c in name_candidates if c in lower), None)
        if email_hit and name_hit:
            email_col_index = lower.index(email_hit)
            fname_col_index = lower.index(name_hit)
            header_row_index = r_idx + 1
            logging.info(f"Header row found at row {header_row_index} (email='{email_hit}' col={email_col_index}, name='{name_hit}' col={fname_col_index})")
            break

    if header_row_index == -1:
        logging.error(f"Could not find header row with email and name columns in the first {MAX_HEADER_SEARCH_ROWS} rows")
        return []

    email_regex = r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$"
    seen_emails = set()

    for row_idx, row in enumerate(
        sheet.iter_rows(min_row=header_row_index + 1, values_only=True),
        start=header_row_index + 1,
    ):
        try:
            email_raw = row[email_col_index]
            fname_raw = row[fname_col_index]

            email = str(email_raw).strip() if email_raw else None
            fname = str(fname_raw).strip() if fname_raw else ""

            if not email or not re.match(email_regex, email):
                logging.debug(f"Row {row_idx}: skipping invalid email '{email_raw}'")
                continue

            if email.lower() in seen_emails:
                logging.debug(f"Row {row_idx}: skipping duplicate email '{email}'")
                continue

            if email.lower() in PINNED_EMAILS:
                logging.debug(f"Row {row_idx}: skipping pinned email '{email}'")
                seen_emails.add(email.lower())
                continue

            if not fname:
                logging.warning(f"Row {row_idx}: missing first name for '{email}', defaulting to empty string")

            contacts.append({
                "email": email,
                "attributes": {"FIRSTNAME": fname},
                "listIds": [target_list_id],
                "updateEnabled": True,
            })
            seen_emails.add(email.lower())

        except IndexError:
            logging.warning(f"Row {row_idx}: fewer columns than expected, skipping")
        except Exception as e:
            logging.exception(f"Row {row_idx}: unexpected error — {e}")

    logging.info(f"Parsed {len(contacts)} unique contacts from {filepath.name} (pinned emails excluded)")
    return contacts


def main():
    setup_logging()
    logging.info("=== Brevo Batch Campaign Runner ===")

    all_contacts = parse_samlet_xlsx(EXCEL_FILE, TARGET_LIST_ID)
    if not all_contacts:
        logging.error("No valid contacts found in Excel file. Aborting.")
        return

    total = len(all_contacts)
    state = load_state()
    offset = state.get("offset", 0)

    if offset >= total:
        logging.warning(f"Offset {offset} >= total {total}. Resetting to 0.")
        offset = 0

    brevo = BrevoRequester()

    logging.info("--- Step 1: Deleting all existing contacts ---")
    deleted, failed = brevo.delete_all_contacts_in_list(TARGET_LIST_ID)
    logging.info(f"Deletion result: {deleted} deleted, {failed} failed")

    logging.info(f"--- Step 2: Uploading batch starting at offset {offset} ---")
    batch_slice = all_contacts[offset: offset + BATCH_SIZE]
    batch = PINNED_CONTACTS + batch_slice
    batch_end = offset + len(batch_slice)

    logging.info(f"Uploading {len(batch)} contacts ({len(PINNED_CONTACTS)} pinned + {len(batch_slice)} from list rows {offset + 1}–{batch_end})")
    brevo.update_all_members_contact_list(batch, target_list_id=TARGET_LIST_ID)

    new_offset = batch_end
    if new_offset >= total:
        logging.info(f"Cycle complete — all {total} contacts have been batched. Resetting offset to 0.")
        new_offset = 0
    else:
        logging.info(f"Batch done. Next run will start at offset {new_offset} (contact {new_offset + 1} of {total}).")

    save_state(new_offset, total)
    logging.info("=== Done ===")


if __name__ == "__main__":
    main()
